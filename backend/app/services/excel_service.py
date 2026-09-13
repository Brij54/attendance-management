import os,re,tempfile
from datetime import date,datetime
from copy import copy
import msoffcrypto
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

def norm(v):
    if v is None:return ""
    if isinstance(v,float) and v.is_integer():v=int(v)
    return str(v).strip().upper()

def parse_date(v):
    if v is None:return None
    if isinstance(v,datetime):return v.date()
    if isinstance(v,date):return v
    s=re.sub(r"\\s+"," ",str(v).strip())
    for f in ["%d %b %y","%d %b %Y","%d/%m/%Y","%d-%m-%Y","%d/%m/%y","%d-%m-%y","%Y-%m-%d","%Y/%m/%d","%d %B %y","%d %B %Y"]:
        try:return datetime.strptime(s,f).date()
        except ValueError:pass
    return None

def parse_swipe(v):
    if v is None:return None
    if isinstance(v,datetime):return v
    if isinstance(v,date):return datetime.combine(v,datetime.min.time())
    s=re.sub(r"\\s+"," ",str(v).strip())
    fmts=["%d/%m/%Y %H:%M:%S","%d/%m/%Y %H:%M","%d/%m/%Y %I:%M:%S %p","%d/%m/%Y %I:%M %p","%d-%m-%Y %H:%M:%S","%d-%m-%Y %H:%M","%d-%m-%Y %I:%M:%S %p","%d-%m-%Y %I:%M %p","%Y-%m-%d %H:%M:%S","%Y-%m-%d %H:%M","%d %b %Y %H:%M:%S","%d %b %Y %H:%M","%d %b %Y %I:%M:%S %p","%d %b %Y %I:%M %p"]
    for f in fmts:
        try:return datetime.strptime(s,f)
        except ValueError:pass
    try:return datetime.fromisoformat(s)
    except ValueError:return None

def header(ws,name):
    t=name.lower()
    for row in ws.iter_rows():
        for c in row:
            if c.value is not None and str(c.value).strip().lower()==t:return c.column,c.row
    return None,None

def decrypt(path,password):
    fd,tmp=tempfile.mkstemp(suffix=".xlsx");os.close(fd)
    try:
        with open(path,"rb") as f:
            o=msoffcrypto.OfficeFile(f);o.load_key(password=password)
            with open(tmp,"wb") as out:o.decrypt(out)
        return tmp
    except Exception as e:
        if os.path.exists(tmp):os.remove(tmp)
        raise ValueError("Could not decrypt Excel file. Check the password.") from e

def open_book(path,password):
    try:return load_workbook(path,data_only=False),None
    except Exception:
        tmp=decrypt(path,password)
        try:return load_workbook(tmp,data_only=False),tmp
        except Exception as e:
            os.remove(tmp);raise ValueError("Decrypted Excel could not be read.") from e

def read_swipes(path, password):
    """Read swipe log. Returns (data, valid, invalid, emp_names).
    emp_names = {emp_code: emp_name} extracted from Employee Name column.
    """
    wb, tmp = open_book(path, password)
    try:
        ws = wb.active
        cc, hr1 = header(ws, "Employee Code")
        sc, hr2 = header(ws, "Swipes")
        if cc is None: raise ValueError("'Employee Code' not found in Swipe Excel.")
        if sc is None: raise ValueError("'Swipes' not found in Swipe Excel.")

        # Find Employee Name column on the header row
        hr = max(hr1 or 1, hr2 or 1)
        nc = None
        name_variants = {"employee name", "emp name", "name", "staff name", "employee_name"}
        for c in range(1, ws.max_column + 1):
            val = str(ws.cell(hr, c).value or "").strip().lower()
            if val in name_variants:
                nc = c
                break

        data = {}; emp_names = {}; valid = invalid = 0
        start_row = hr + 1
        for r in range(start_row, ws.max_row + 1):
            code = norm(ws.cell(r, cc).value)
            raw  = ws.cell(r, sc).value
            if not code: continue
            # Collect name whenever we see this employee
            if nc and code not in emp_names:
                n = str(ws.cell(r, nc).value or "").strip()
                if n:
                    emp_names[code] = n
            if raw in (None, ""): continue
            dt = parse_swipe(raw)
            if not dt: invalid += 1; continue
            valid += 1; data.setdefault(code, {})
            old = data[code].get(dt.date())
            if old is None or dt < old: data[code][dt.date()] = dt
        return data, valid, invalid, emp_names
    finally:
        wb.close()
        if tmp and os.path.exists(tmp): os.remove(tmp)

def update_attendance(path,password,data,outpath):
    wb,tmp=open_book(path,password)
    try:
        ws=wb.active;cc,hr=header(ws,"Employee Code")
        if cc is None:raise ValueError("'Employee Code' not found in Attendance Excel.")
        dates=[(c,parse_date(ws.cell(hr,c).value)) for c in range(1,ws.max_column+1)]
        dates=[x for x in dates if x[1]]
        if not dates:raise ValueError("No date columns such as '21 Jul 26' found.")

        # Unmerge top title banners (rows above header row) to prevent openpyxl text shifting/duplication
        title_merged_ranges = []
        if hr and hr > 1:
            merged_list = list(ws.merged_cells.ranges)
            for rng in merged_list:
                if rng.min_row < hr:
                    title_merged_ranges.append((rng.min_row, rng.max_row, rng.min_col, rng.max_col))
                    ws.unmerge_cells(range_string=str(rng))

        for c,d in reversed(dates):
            ws.insert_cols(c+1,1);pc=c+1
            src=ws.cell(hr,c);dst=ws.cell(hr,pc)
            dst.value=d.strftime("%d-%b")+" Punchin Time"

            # Clear title row cells for inserted columns
            if hr and hr > 1:
                for r_top in range(1, hr):
                    ws.cell(r_top, pc).value = None

            if src.has_style:
                dst.font=copy(src.font);dst.fill=copy(src.fill);dst.border=copy(src.border);dst.alignment=copy(src.alignment);dst.protection=copy(src.protection)
            ws.column_dimensions[get_column_letter(pc)].width=max(ws.column_dimensions[get_column_letter(c)].width or 12,21)
            for r in range(hr+1,ws.max_row+1):
                s=ws.cell(r,c);x=ws.cell(r,pc)
                if s.has_style:
                    x.font=copy(s.font);x.fill=copy(s.fill);x.border=copy(s.border);x.alignment=copy(s.alignment);x.protection=copy(s.protection)
                code=norm(ws.cell(r,cc).value);dt=data.get(code,{}).get(d)
                x.value=dt
                if dt:x.number_format="dd/mm/yyyy hh:mm:ss"

        # Re-merge title banners across full expanded width
        num_inserted = len(dates)
        for min_r, max_r, min_c, max_c in title_merged_ranges:
            ws.merge_cells(start_row=min_r, end_row=max_r, start_column=min_c, end_column=max_c + num_inserted)

        wb.save(outpath)
        return len(dates)
    finally:
        wb.close()
        if tmp and os.path.exists(tmp):os.remove(tmp)


def read_attendance_employees(path, password):
    """Extract {emp_code: emp_name} mapping from Attendance Master Excel.
    Tries common Employee Name column header variants.
    """
    wb, tmp = open_book(path, password)
    try:
        ws = wb.active
        cc, hr = header(ws, "Employee Code")
        if cc is None:
            return {}
        # Try multiple possible column names for employee name
        nc = None
        name_variants = {"employee name", "emp name", "name", "staff name", "employee_name"}
        for c in range(1, ws.max_column + 1):
            val = str(ws.cell(hr, c).value or "").strip().lower()
            if val in name_variants:
                nc = c
                break
        emp_map = {}
        for r in range(hr + 1, ws.max_row + 1):
            code = norm(ws.cell(r, cc).value)
            if not code:
                continue
            name = str(ws.cell(r, nc).value or "").strip() if nc else ""
            emp_map[code] = name
        return emp_map
    finally:
        wb.close()
        if tmp and os.path.exists(tmp):
            os.remove(tmp)
